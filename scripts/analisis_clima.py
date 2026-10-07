"""Análisis climático para epidemiología del Mal Suramericano de las Hojas (SALB) y de
Colletotrichum Leaf Fall (CLF) en caucho.

Sitios: Finca Las Delicias y C.I. Taluma - AGROSAVIA (estación Las Margaritas a 0.9 km).
Ventana epidemiológica continua: 1 nov - 28/29 feb.

Ejecutar:  python scripts/descargar_chirps.py     (una vez: lluvia CHIRPS)
           python scripts/descargar_ideam_aut.py  (una vez: HR horaria La Palomera y La Libertad)
           python scripts/analisis_clima.py
"""
import numpy as np
import pandas as pd

import config as C
import graficos
import humedad
from cargar_datos import cargar_sitio, leer_agera5, leer_estaciones, leer_imerg

FUENTES = {"era5": "ERA5-Land / CHIRPS", "nasa": "NASA POWER", "agera5": "AgERA5",
           "imerg": "GPM IMERG", "est": "Estación Las Margaritas", "ref": "Estación ref. HR"}
VARIABLES = {"tmax": "T máxima (°C)", "tmin": "T mínima (°C)", "tmed": "T media (°C)",
             "prec": "Precipitación (mm)", "hr": "Humedad relativa (%)",
             "rad": "Radiación (MJ/m²/día)"}


# ---------------------------------------------------------------- indicadores
def indicadores_horarios(h, sufijo):
    """Horas de HR alta (proxy de hoja mojada) y condiciones horarias para SALB."""
    dia = h.index.normalize()
    d = pd.DataFrame({f"hr90_horas_{sufijo}": humedad.horas_hr_alta(h["hr"])})
    if "t" in h:
        cond = (h["hr"] >= C.SALB["hr_umbral"]) & h["t"].between(C.SALB["t_min"], C.SALB["t_max"])
        d[f"salb_horas_cond_{sufijo}"] = cond.groupby(dia).sum().where(d[f"hr90_horas_{sufijo}"].notna())
        noche = h[(h.index.hour >= 18) | (h.index.hour < 6)]
        d[f"t_noche_{sufijo}"] = noche["t"].groupby(noche.index.normalize()).mean()
        # Déficit de presión de vapor (kPa): < 0.5 kPa = aire casi saturado, favorable a hongos
        es = 0.6108 * np.exp(17.27 * h["t"] / (h["t"] + 237.3))
        d[f"dpv_{sufijo}"] = (es * (1 - h["hr"] / 100)).groupby(dia).mean()
    n = h["hr"].groupby(dia).count()
    d[f"hr_{sufijo}"] = h["hr"].groupby(dia).mean().where(n >= 20)
    return d


def dias_favorables(hr90, tmed, lluvia):
    salb = ((hr90 >= C.SALB["horas_min"]) & tmed.between(C.SALB["t_min"], C.SALB["t_max"]))
    humedo = (hr90 >= C.CLF["horas_min"]) | (lluvia >= C.CLF["lluvia_mm"])
    clf = humedo & tmed.between(C.CLF["t_min"], C.CLF["t_max"])
    validos = hr90.notna() & tmed.notna()
    return salb.astype(float).where(validos), clf.astype(float).where(validos)


def indicadores_epidemiologicos(df):
    """Días favorables (1/0). Principal: HR horaria ERA5-Land corregida (QM con la estación de
    referencia), T ERA5-Land y lluvia CHIRPS. Sensibilidad: otras fuentes de HR."""
    t, lluvia = df["tmed_era5"], df["prec_era5"]
    df["dia_lluvia"] = (lluvia >= C.LLUVIA_DIA_MM).astype(float).where(lluvia.notna())
    df["hr90_horas"] = df["hr90_horas_era5c"]
    df["salb_favorable"], df["clf_favorable"] = dias_favorables(df["hr90_horas_era5c"], t, lluvia)
    for suf in ["nasa", "era5", "alt", "ref"]:
        df[f"salb_fav_{suf}"], df[f"clf_fav_{suf}"] = dias_favorables(df[f"hr90_horas_{suf}"], t, lluvia)
    grupo = (df["dia_lluvia"] != df["dia_lluvia"].shift()).cumsum()
    df["racha_lluvia"] = df["dia_lluvia"].groupby(grupo).cumsum() * df["dia_lluvia"]
    return df


def ventanas_moviles(df):
    """Acumulados / conteos en los N días previos (incluye el día). Para relacionar con
    las fechas de evaluación de incidencia/severidad en campo."""
    v = pd.DataFrame(index=df.index)
    for n in C.VENTANAS:
        r = df.rolling(n, min_periods=n)
        v[f"prec_{n}d"] = r["prec_era5"].sum()
        v[f"dias_lluvia_{n}d"] = r["dia_lluvia"].sum()
        v[f"hr90_horas_{n}d"] = r["hr90_horas"].sum()
        v[f"tmed_{n}d"] = r["tmed_era5"].mean()
        v[f"salb_fav_{n}d"] = r["salb_favorable"].sum()
        v[f"clf_fav_{n}d"] = r["clf_favorable"].sum()
    return v


# ---------------------------------------------------------------- medias móviles
def media_movil(df, cols, lluvia):
    """Para cada N en C.MA_DIAS:
      _ma{N}c  centrada (N días alrededor del día)  -> describir y graficar
      _ma{N}r  retrasada (el día y los N-1 previos)  -> alertas y relación con la enfermedad
    La lluvia no se promedia: se acumula (_acum{N}c / _acum{N}r)."""
    out = pd.DataFrame(index=df.index)
    for n in C.MA_DIAS:
        for c in cols:
            out[f"{c}_ma{n}c"] = df[c].rolling(n, center=True, min_periods=n).mean()
            out[f"{c}_ma{n}r"] = df[c].rolling(n, min_periods=n).mean()
        for c in lluvia:
            out[f"{c}_acum{n}c"] = df[c].rolling(n, center=True, min_periods=n).sum()
            out[f"{c}_acum{n}r"] = df[c].rolling(n, min_periods=n).sum()
    return out


def efecto_suavizado(df, cols):
    """Qué tan 'errática' es cada serie: desviación del cambio día a día y autocorrelación,
    original y con cada media móvil."""
    filas = []
    for c in cols:
        s = df[c].dropna()
        if len(s) < 60:
            continue
        fila = {"serie": c, "desv_cambio_diario_original": s.diff().std(), "autocorr_lag1_original": s.autocorr(1)}
        for n in C.MA_DIAS:
            for tipo, centro in [("c", True), ("r", False)]:
                m = df[c].rolling(n, center=centro, min_periods=n).mean().dropna()
                fila[f"desv_cambio_diario_ma{n}{tipo}"] = m.diff().std()
                fila[f"reduccion_ruido_ma{n}_%"] = 100 * (1 - m.diff().std() / s.diff().std())
                fila[f"autocorr_lag1_ma{n}{tipo}"] = m.autocorr(1)
        filas.append(fila)
    return pd.DataFrame(filas).round(3)


# ---------------------------------------------------------------- agregaciones
AGG = {"prec": "sum", "tmax": "mean", "tmin": "mean", "tmed": "mean", "hr": "mean",
       "rad": "mean", "etp": "sum", "bs": "sum", "hr90_horas": "sum", "salb_horas_cond": "sum",
       "dpv": "mean", "t_noche": "mean", "dia_lluvia": "sum", "salb_fav": "sum",
       "clf_fav": "sum", "racha_lluvia": "max", "horas_lluvia": "sum", "intensidad_max": "max"}


def regla(col):
    for k in sorted(AGG, key=len, reverse=True):  # el prefijo más largo primero (hr90_horas antes que hr)
        if col.startswith(k):
            return AGG[k]
    return "mean"


def agregar(df, freq, min_frac=0.8):
    """Agrega a semana/mes. Solo si hay >= 80 % de días con dato (evita totales falsos)."""
    g = df.resample(freq)
    res = g.agg({c: regla(c) for c in df.columns})
    dias = g.size()
    res = res.where(g.count().div(dias, axis=0) >= min_frac)
    res.insert(0, "dias", dias)
    return res


def temporada(idx):
    """Temporada epidemiológica: 1 nov (año) - 28/29 feb (año+1). Fuera de la ventana = ''."""
    idx = pd.DatetimeIndex(idx)
    ini = np.where(idx.month >= 11, idx.year, idx.year - 1)
    lab = pd.Series([f"{a}-{a + 1}" for a in ini], index=idx)
    return lab.where(idx.month.isin(C.VENTANA_MESES), "").values


def semanal_calendario(df):
    """Semana calendario ISO (lunes a domingo)."""
    s = agregar(df, "W-SUN")
    s.index = s.index - pd.Timedelta(days=6)  # etiqueta = lunes de inicio
    s.index.name = "inicio_semana"
    iso = s.index.isocalendar()
    s.insert(0, "anio_iso", iso.year.values)
    s.insert(1, "semana_iso", iso.week.values)
    # Semana dentro de la ventana si el lunes de inicio cae entre 1 nov y 28/29 feb
    s.insert(2, "temporada", temporada(s.index))
    return s


# ---------------------------------------------------------------- validación
def metricas(obs, sim):
    m = pd.concat([obs, sim], axis=1).dropna()
    if len(m) < 10:
        return None
    o, s = m.iloc[:, 0], m.iloc[:, 1]
    return {"n": len(m), "media_ref": o.mean(), "media_prod": s.mean(), "sesgo": (s - o).mean(),
            "MAE": (s - o).abs().mean(), "RMSE": np.sqrt(((s - o) ** 2).mean()),
            "r": o.corr(s), "r_spearman": o.corr(s, method="spearman")}


def _suma_completa(o, s, freq):
    m = pd.concat([o, s], axis=1).dropna()
    g = m.resample(freq)
    ok = g.size() >= (6 if freq == "W-SUN" else 25)
    tot = g.sum()[ok]
    return tot.iloc[:, 0], tot.iloc[:, 1]


def validar_taluma(d):
    """Productos en Taluma vs estación Las Margaritas (0.9 km)."""
    filas = []
    prods = [f for f in ("era5", "nasa", "agera5", "imerg") if any(c.endswith("_" + f) for c in d)]
    for var in ["tmax", "tmin", "tmed", "prec"]:
        for prod in prods:
            if f"{var}_{prod}" not in d:
                continue
            o, s = d[f"{var}_est"], d[f"{var}_{prod}"]
            escalas = [("diaria", o, s)]
            if var == "prec":
                escalas += [("semanal", *_suma_completa(o, s, "W-SUN")), ("mensual", *_suma_completa(o, s, "MS"))]
            for esc, oo, ss in escalas:
                r = metricas(oo, ss)
                if r:
                    filas.append({"variable": var, "producto": FUENTES[prod], "escala": esc, **r})
    return pd.DataFrame(filas).round(3)


def deteccion_lluvia(d):
    """Contingencia día lluvioso (>= 1 mm): POD, FAR, CSI y sesgo de frecuencia."""
    filas = []
    for prod in [f for f in ("era5", "nasa", "agera5", "imerg") if f"prec_{f}" in d]:
        m = d[["prec_est", f"prec_{prod}"]].dropna() >= C.LLUVIA_DIA_MM
        o, s = m.iloc[:, 0], m.iloc[:, 1]
        a, b, c = (o & s).sum(), (~o & s).sum(), (o & ~s).sum()
        filas.append({"producto": FUENTES[prod], "n": len(m), "aciertos": a, "falsas_alarmas": b,
                      "omisiones": c, "POD": a / (a + c), "FAR": b / (a + b), "CSI": a / (a + b + c),
                      "sesgo_frecuencia": (a + b) / (a + c)})
    return pd.DataFrame(filas).round(3)


def desfase_lluvia(d):
    """El 'día pluviométrico' IDEAM va de 7 a.m. a 7 a.m.: ¿conviene desplazar 1 día?"""
    filas = []
    for prod in [f for f in ("era5", "nasa", "agera5", "imerg") if f"prec_{f}" in d]:
        for lag in (-1, 0, 1):
            filas.append({"producto": FUENTES[prod], "desfase_dias": lag,
                          "r_diaria": d["prec_est"].corr(d[f"prec_{prod}"].shift(lag))})
    return pd.DataFrame(filas).round(3)


def validar_agera5_hr():
    """HR de AgERA5 (promedio 06-18 h) vs estaciones de referencia, si se exportó desde GEE."""
    filas = []
    for ref, info in C.REF_HR.items():
        ag = leer_agera5(ref)
        if ag is None or "hr_agera5" not in ag:
            continue
        o = humedad.leer_estacion_horaria(ref)["hr"]
        o = o[o.index.hour.isin([6, 9, 12, 15, 18])]
        o = o.groupby(o.index.normalize()).mean()
        r = metricas(o, ag["hr_agera5"])
        if r:
            filas.append({"estacion": info["nombre"], "producto": "AgERA5 HR media 06-18 h", **r})
    return pd.DataFrame(filas).round(3)


def comparar_sitios(sitios):
    """Las Delicias vs Taluma con las mismas fuentes: diferencias medias y sincronía."""
    a, b = sitios["delicias"], sitios["taluma"]
    filas = []
    for c in ["tmax_era5", "tmin_era5", "tmed_era5", "hr_era5c", "prec_era5", "hr90_horas",
              "salb_favorable", "clf_favorable", "hr90_horas_ref", "tmed_nasa", "hr_nasa", "prec_nasa"]:
        for esc, fa, fb in [("diaria", a["diario"][c], b["diario"][c]),
                            ("mensual", a["mensual"][c], b["mensual"][c])]:
            m = pd.concat([fa, fb], axis=1).dropna()
            if len(m) < 10:
                continue
            filas.append({"variable": c, "escala": esc, "media_delicias": m.iloc[:, 0].mean(),
                          "media_taluma": m.iloc[:, 1].mean(),
                          "diferencia_taluma_menos_delicias": (m.iloc[:, 1] - m.iloc[:, 0]).mean(),
                          "r_entre_sitios": m.iloc[:, 0].corr(m.iloc[:, 1])})
    return pd.DataFrame(filas).round(3)


# ---------------------------------------------------------------- climatología y temporadas
def climatologia_mensual(mensual):
    cols = [c for c in mensual.columns if c != "dias"]
    clim = mensual[cols].groupby(mensual.index.month.rename("mes")).agg(["mean", "std"]).round(2)
    clim.columns = [f"{a}_{b}" for a, b in clim.columns]
    clim.insert(0, "mes_nombre", [C.NOMBRE_MES[i] for i in clim.index])
    clim.insert(1, "en_ventana", ["Sí" if i in C.VENTANA_MESES else "" for i in clim.index])
    return clim


def anomalias_ventana(mensual, cols):
    m = mensual[cols]
    z = ((m - m.groupby(m.index.month).transform("mean")) / m.groupby(m.index.month).transform("std"))
    z = z.add_suffix("_anom_z").round(2)
    z.insert(0, "mes", m.index.month)
    return z[z["mes"].isin(C.VENTANA_MESES)]


def resumen_temporadas(diario):
    d = diario[diario.index.month.isin(C.VENTANA_MESES)].copy()
    d["temporada"] = temporada(d.index)
    d["mes"] = d.index.month
    agg = {"prec_era5": "sum", "dia_lluvia": "sum", "tmed_era5": "mean", "tmin_era5": "mean",
           "tmax_era5": "mean", "hr_era5c": "mean", "hr90_horas": "mean", "dpv_era5c": "mean",
           "salb_favorable": "sum", "clf_favorable": "sum", "salb_fav_nasa": "sum", "clf_fav_nasa": "sum",
           "salb_fav_era5": "sum", "salb_fav_alt": "sum", "hr90_horas_ref": "mean", "salb_fav_ref": "sum"}
    por_mes = d.groupby(["temporada", "mes"]).agg(agg).round(2)
    por_temp = d.groupby("temporada").agg(agg).round(2)
    por_temp["dias_con_hr_ref"] = d.groupby("temporada")["hr90_horas_ref"].count()
    por_temp = por_temp[d.groupby("temporada").size() >= 118]  # temporadas completas (nov-feb)
    return por_temp, por_mes.reset_index()


# ---------------------------------------------------------------- por sitio
def procesar_sitio(sitio, est):
    info = C.SITIOS[sitio]
    ref = info["ref_hr"]
    otra = next(r for r in C.REF_HR if r != ref)
    print(f"\n=== {info['nombre']} (HR de referencia: {C.REF_HR[ref]['nombre']}) ===")
    diario, h_era5, h_nasa = cargar_sitio(sitio, est)
    diario = diario.rename(columns={"hr_era5": "hr_era5_sincorr"})
    h_era5c, est_ref = humedad.corregir_sitio(sitio, h_era5)
    tabla_alt = humedad.entrenar_qm(humedad.leer_estacion_horaria(otra)["hr"],
                                    humedad.leer_era5_horario(otra)["hr"])
    h_alt = h_era5.assign(hr=humedad.aplicar_qm(h_era5["hr"], tabla_alt))
    for h, suf in [(h_era5c, "era5c"), (h_era5, "era5"), (h_alt, "alt"), (h_nasa, "nasa"), (est_ref, "ref")]:
        diario = diario.join(indicadores_horarios(h, suf).drop(columns=["hr_era5", "hr_nasa"], errors="ignore"),
                             how="left")
    diario = indicadores_epidemiologicos(diario)

    cols_ma = [c for c in diario if c.split("_")[0] in ("tmax", "tmin", "tmed", "hr", "rad")
               and not c.startswith("hr90") and diario[c].notna().sum() > 60] + \
              ["hr90_horas", "hr90_horas_nasa", "hr90_horas_ref", "dpv_era5c", "t_noche_era5c"]
    lluvia = [c for c in diario if c.startswith("prec_")]
    ma = media_movil(diario, cols_ma, lluvia)
    mensual = agregar(diario, "MS")
    mensual.index.name = "mes"
    res = dict(diario=diario, ma=ma, mensual=mensual, suav=efecto_suavizado(diario, cols_ma),
               vent=ventanas_moviles(diario), semanal=semanal_calendario(diario),
               clim=climatologia_mensual(mensual),
               anom=anomalias_ventana(mensual, ["prec_era5", "tmed_era5", "hr_era5c", "hr90_horas",
                                                "salb_favorable", "clf_favorable"]))
    res["temp"], res["temp_mes"] = resumen_temporadas(diario)
    if info["estacion"]:
        res["valid"] = validar_taluma(diario)
        res["conting"] = deteccion_lluvia(diario)
        res["desfase"] = desfase_lluvia(diario)

    carpeta = C.SALIDA / sitio
    carpeta.mkdir(parents=True, exist_ok=True)
    ruta = carpeta / f"Base_climatica_{sitio}.xlsx"
    hojas = {"diario": diario.round(3), "medias_moviles": ma.round(3), "ventanas_moviles": res["vent"].round(2),
             "semanal_calendario": res["semanal"].round(2),
             "semanal_ventana": res["semanal"][res["semanal"]["temporada"] != ""].round(2),
             "mensual": mensual.round(2), "climatologia_mensual": res["clim"],
             "anomalias_ventana": res["anom"], "resumen_temporadas": res["temp"],
             "temporada_por_mes": res["temp_mes"], "efecto_medias_moviles": res["suav"]}
    if info["estacion"]:
        hojas.update({"validacion_estacion": res["valid"], "deteccion_lluvia": res["conting"],
                      "desfase_lluvia_estacion": res["desfase"]})
    escribir_excel(ruta, hojas, leeme(sitio))
    graficos.por_sitio(sitio, res)
    print(f"   -> {ruta}")
    return res


def escribir_excel(ruta, hojas, notas):
    with pd.ExcelWriter(ruta, engine="openpyxl") as xw:
        pd.DataFrame({"Hoja / dato": list(notas), "Descripción": list(notas.values())}).to_excel(
            xw, sheet_name="LEEME", index=False)
        for nombre, df in hojas.items():
            if df is not None and len(df):
                df.to_excel(xw, sheet_name=nombre, index=not isinstance(df.index, pd.RangeIndex))
    _formato(ruta)


def leeme(sitio):
    info = C.SITIOS[sitio]
    ref = C.REF_HR[info["ref_hr"]]
    otra = next(v for k, v in C.REF_HR.items() if k != info["ref_hr"])
    n = {"Sitio": f"{info['nombre']} ({info['municipio']}, Meta) lat {info['lat']}, lon {info['lon']}",
         "Ventana epidemiológica": "Continua del 1 de noviembre al 28/29 de febrero (columna 'temporada').",
         "Enfermedades": f"{C.SALB_NOMBRE}, Pseudocercospora ulei; {C.CLF_NOMBRE}, Colletotrichum spp.",
         "Sufijo _era5": "T: ERA5-Land (9 km). Radiación y ETP: ERA5 (0.25°). Lluvia: CHIRPS v2.0 (5 km).",
         "Sufijo _era5c": f"ERA5-Land horaria con HR corregida por mapeo de cuantiles (mes x bloque de 3 h) "
                          f"contra {ref['nombre']} ({ref['km'][sitio]} km). Fuente PRINCIPAL de HR.",
         "Sufijo _alt": f"ERA5-Land corregida con la otra estación ({otra['nombre']}): sensibilidad.",
         "Sufijo _ref": f"Observado en {ref['nombre']}, a {ref['km'][sitio]} km (HR horaria IDEAM, datos.gov.co).",
         "hr_era5_sincorr": "HR diaria de ERA5-Land sin corregir.",
         "Sufijo _nasa": "NASA POWER (MERRA-2 / CERES, celda ~50 km)."}
    if info["estacion"]:
        n["Sufijo _est"] = "IDEAM Hacienda Las Margaritas [35125010], a 0.9 km: T, lluvia, brillo solar (no mide HR)."
    else:
        n["Estación T / lluvia"] = ("No hay estación climatológica en la finca. Hay pluviómetros IDEAM cercanos "
                                    "(Santa Helena 2.1 km, El Toro 7.3 km, Yaguarito 7.4 km): ver GUIA_ANALISIS.md.")
    n.update({
        "Sufijo _agera5 / _imerg": "AgERA5 y GPM IMERG si se exportaron desde Earth Engine (datos_crudos/gee/).",
        "hr90_horas": "Horas/día con HR >= 90 % (ERA5-Land corregida). Proxy de horas de hoja mojada.",
        "salb_horas_cond_*": "Horas con HR >= 90 % y T entre 20-28 °C.",
        "dpv_*": "Déficit de presión de vapor medio diario (kPa). < 0.5 kPa = aire casi saturado.",
        "t_noche_*": "Temperatura media 18:00-06:00.",
        "salb_favorable": f"1 si hr90_horas >= {C.SALB['horas_min']} y T media ERA5-Land "
                          f"{C.SALB['t_min']}-{C.SALB['t_max']} °C.",
        "clf_favorable": f"Colletotrichum Leaf Fall: 1 si (hr90_horas >= {C.CLF['horas_min']} o lluvia >= "
                         f"{C.CLF['lluvia_mm']} mm) y T media ERA5-Land {C.CLF['t_min']}-{C.CLF['t_max']} °C.",
        "*_fav_nasa / _era5 / _alt / _ref": "Días favorables con otras fuentes de HR (sensibilidad).",
        "dia_lluvia / racha_lluvia": f"Día con CHIRPS >= {C.LLUVIA_DIA_MM} mm / días lluviosos consecutivos.",
        "_ma3c / _ma7c": "Media móvil CENTRADA de 3 / 7 días (días antes y después). Para describir la tendencia.",
        "_ma3r / _ma7r": "Media móvil RETRASADA de 3 / 7 días (el día y los 2 / 6 anteriores). Para alertas "
                         "y para relacionar con la enfermedad (solo usa el pasado).",
        "_acum3r / _acum7r": "Lluvia acumulada en 3 / 7 días (la lluvia se acumula, no se promedia).",
        "ventanas_moviles": f"Acumulados previos de {C.VENTANAS} días. Relacionar con fechas de evaluación.",
        "semanal_calendario": "Semana ISO lunes-domingo. Agregados solo con >= 80 % de días con dato.",
        "anomalias": "z = (valor - media del mes) / desviación. |z| > 1 = mes atípico.",
        "efecto_medias_moviles": "Reducción de la variación día a día (%) y autocorrelación con MA3 y MA7.",
    })
    if info["estacion"]:
        n["validacion_estacion"] = "sesgo = producto - estación. r = correlación de Pearson."
        n["deteccion_lluvia"] = "POD = prob. de detección; FAR = falsas alarmas; CSI = índice de éxito crítico."
    return n


def _formato(ruta):
    from openpyxl import load_workbook
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    wb = load_workbook(ruta)
    for ws in wb.worksheets:
        ws.freeze_panes = "B2"
        for col in ws.columns:
            largo = max(len(str(c.value)) if c.value is not None else 0 for c in col[:50])
            ws.column_dimensions[col[0].column_letter].width = min(max(10, largo + 2), 70)
        # Cursiva para los nombres científicos en los textos
        for fila in ws.iter_rows():
            for c in fila:
                if isinstance(c.value, str) and ("Colletotrichum" in c.value or "Pseudocercospora" in c.value):
                    c.value = _cursiva(c.value, CellRichText, TextBlock, InlineFont)
    wb.save(ruta)


def _cursiva(texto, CellRichText, TextBlock, InlineFont):
    import re
    partes = re.split(r"(Colletotrichum|Pseudocercospora ulei)", texto)
    return CellRichText(*[TextBlock(InlineFont(i=True), p) if p in ("Colletotrichum", "Pseudocercospora ulei") else p
                          for p in partes if p])


# ---------------------------------------------------------------- principal
def main():
    C.SALIDA.mkdir(exist_ok=True)
    print("Leyendo estaciones IDEAM...")
    est, estaciones = leer_estaciones()
    print("Validando horas con HR >= 90 % contra estaciones automáticas...")
    valid_hr, valid_hr_mes = humedad.validar_horas_hr(list(C.REF_HR))

    sitios = {s: procesar_sitio(s, est) for s in C.SITIOS}

    print("\n=== Comparación entre sitios ===")
    temporadas = pd.concat({C.SITIOS[s]["nombre"]: r["temp"] for s, r in sitios.items()}, names=["sitio", "temporada"])
    clim = pd.concat({C.SITIOS[s]["nombre"]: r["clim"] for s, r in sitios.items()}, names=["sitio", "mes"])
    distancias = pd.DataFrame([
        ("Las Delicias", "C.I. Taluma", 138.2), ("C.I. Taluma", "Est. Las Margaritas (T, lluvia)", 0.9),
        ("C.I. Taluma", "Est. La Palomera AUT (HR)", 39.3), ("Las Delicias", "Est. La Libertad AUT (HR)", 30.2),
        ("Las Delicias", "Pluviómetro Santa Helena [35010110]", 2.1),
        ("Las Delicias", "Pluviómetro El Toro [35010060]", 7.3),
        ("Las Delicias", "Pluviómetro Yaguarito [35010150]", 7.4),
        ("Las Delicias", "Climatológica Morichal [3502500006]", 16.9)], columns=["desde", "hasta", "km"])
    notas = {
        "Objetivo": "Comparar Las Delicias y C.I. Taluma con las mismas fuentes en la ventana 1 nov - 28/29 feb.",
        "validacion_Taluma": "Productos en Taluma vs estación Las Margaritas (0.9 km): T y lluvia.",
        "validacion_horas_HR90": "Horas/día con HR >= 90 % de cada producto vs observadas en La Palomera y "
                                 "La Libertad. 'validación cruzada por años' = entrenado en años impares, "
                                 "probado en pares y viceversa. 'entrenado en ...' = transferido de la otra estación.",
        "acierto_dia_10h_%": "% de días en que producto y estación coinciden en si se supera el umbral SALB de 10 h.",
        "sesgo_h_ventana / r_ventana": "Mismas métricas solo en la ventana nov-feb.",
        "comparacion_sitios": "Diferencia Taluma - Delicias y correlación entre sitios.",
    }
    escribir_excel(C.SALIDA / "Comparacion_sitios.xlsx", {
        "distancias_km": distancias,
        "estaciones_DHIME": pd.DataFrame({"variable": list(estaciones), "estacion": list(estaciones.values())}),
        "validacion_Taluma": sitios["taluma"]["valid"],
        "deteccion_lluvia_Taluma": sitios["taluma"]["conting"],
        "desfase_lluvia_Taluma": sitios["taluma"]["desfase"],
        "validacion_horas_HR90": valid_hr, "horas_HR90_por_mes": valid_hr_mes,
        "validacion_AgERA5_HR": validar_agera5_hr(),
        "comparacion_sitios": comparar_sitios(sitios), "temporadas_ambos": temporadas,
        "climatologia_ambos": clim,
    }, notas)
    graficos.comparacion(sitios, valid_hr, valid_hr_mes)
    print(f"\nListo. Resultados en {C.SALIDA}")
    return sitios


if __name__ == "__main__":
    main()
